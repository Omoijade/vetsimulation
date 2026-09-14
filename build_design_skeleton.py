from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("/Users/thomas-julianirabor/Downloads/SIMULATION/Reusable_Interactive_Design_Skeleton.docx")

NAVY = "18334A"
BLUE = "2E74B5"
TEAL = "2E7D7A"
INK = "1F2933"
MUTED = "5F6C78"
PALE_BLUE = "E8EEF5"
PALE_TEAL = "E8F3F2"
PALE_GOLD = "FFF6DE"
PALE_RED = "FCEBEC"
LINE = "D6DEE5"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_paragraph_callout(paragraph, fill, accent):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "7")
    left.set(qn("w:color"), accent)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "180")
    ind.set(qn("w:right"), "140")
    p_pr.append(ind)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run(run, size=11, color=INK, bold=False, italic=False, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def add_numbering(doc, abstract_id, num_id, kind):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"A1B2C{abstract_id:03d}"[-8:])
    abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "\u2022" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    lvl.append(p_pr)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def add_para(doc, text="", style=None, size=11, color=INK, bold=False, italic=False,
             before=0, after=6, line=1.25, align=WD_ALIGN_PARAGRAPH.LEFT,
             keep=False):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if keep:
        set_keep_with_next(p)
    if text:
        set_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text, bold_lead=None):
    p = add_para(doc, after=4, line=1.25)
    apply_num(p, 41)
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]))
    else:
        set_run(p.add_run(text))
    return p


def add_step(doc, title, detail):
    p = add_para(doc, after=4, line=1.25, keep=True)
    apply_num(p, 42)
    set_run(p.add_run(title), bold=True, color=NAVY)
    set_run(p.add_run(f" - {detail}"))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    set_keep_with_next(p)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, accent=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.2
    set_paragraph_callout(p, fill, accent)
    set_run(p.add_run(f"{label}: "), size=10.5, color=accent, bold=True)
    set_run(p.add_run(text), size=10.5, color=INK)
    return p


def add_definition(doc, term, definition):
    p = add_para(doc, after=5, line=1.2)
    set_run(p.add_run(f"{term}. "), bold=True, color=NAVY)
    set_run(p.add_run(definition))
    return p


def set_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in (
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, NAVY, 10, 5),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    set_run(hp.add_run("INTERACTIVE SYSTEM DESIGN SKELETON"), size=8.5, color=MUTED, bold=True)
    footer = section.footer
    set_page_field(footer.paragraphs[0])

    add_numbering(doc, 31, 41, "bullet")
    add_numbering(doc, 32, 42, "decimal")


def add_cover(doc):
    add_para(doc, "FIELD GUIDE", size=10, color=TEAL, bold=True, after=28, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Interactive System Design Skeleton", size=29, color=NAVY, bold=True, after=10,
             line=1.05, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "A reusable guide for designing decision-centred simulations, sandboxes, and learning tools",
             size=14, color=MUTED, after=34, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Build choices that matter. Make constraints discoverable. Turn actions into consequences that can be seen, compared, explained, and revisited.",
             size=12, color=INK, after=42, line=1.35, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_callout(doc, "Scope", "This guide concerns the design of the experience around subject matter: objectives, models, choices, representations, feedback, interaction trust, and facilitation. It is intentionally domain-neutral.", fill=PALE_TEAL, accent=TEAL)
    add_para(doc, "Version 1.0 | June 2026", size=9.5, color=MUTED, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_intro(doc):
    add_heading(doc, "How to use this guide", 1)
    add_para(doc, "Use the guide before interface design begins, then return to it during prototyping and testing. It is a skeleton: each project must supply its own subject matter, rules, evidence, and constraints.")
    add_callout(doc, "Central test", "At every moment, a participant should be able to tell what is happening, what can be changed, what may happen next, and what actually happened after a decision.", fill=PALE_GOLD, accent="8A6500")
    add_heading(doc, "The four layers", 2)
    add_definition(doc, "System model", "The state, resources, actors, constraints, relationships, and transition rules that make the experience behave consistently.")
    add_definition(doc, "Choice model", "The interventions available to participants, including their costs, limits, trade-offs, timing, reversibility, and dependencies.")
    add_definition(doc, "Representation model", "The views, visual encodings, comparisons, and explanations that make state and consequence interpretable without revealing every answer.")
    add_definition(doc, "Facilitation model", "The rules, prompts, records, and debrief structures that help participants reconstruct why an outcome occurred.")
    add_heading(doc, "Seven-step design sequence", 2)
    for title, detail in (
        ("Define the objective", "Separate the participant's goal from the educational or evaluative purpose."),
        ("Construct the system", "Specify state, actors, resources, constraints, demand, and causal relationships."),
        ("Design the choice space", "Offer several viable paths with visible trade-offs and no single scripted sequence."),
        ("Design the representations", "Give each view a distinct job and connect overview, cause, action, forecast, and history."),
        ("Make consequences concrete", "Show accumulation, delay, interaction effects, and before/after changes."),
        ("Test for trust and clarity", "Eliminate silent state changes, vague terms, dead controls, and unreachable fixes."),
        ("Prepare debrief and iteration", "Capture actions and outcomes in forms that can be discussed, compared, and written about."),
    ):
        add_step(doc, title, detail)


def add_commitments(doc):
    add_heading(doc, "1. Core design commitments", 1)
    commitments = [
        ("Autonomy within a legible objective", "State the destination and boundaries, but do not prescribe the route. Participants need enough options to form, test, and revise their own strategy."),
        ("Scarcity that can be diagnosed", "Constraints should create pressure, not confusion. Participants must be able to infer what is missing from patterns in demand, capacity, cost, quality, access, time, or relationships."),
        ("Every visible element earns its place", "A component should support orientation, comparison, diagnosis, action, prediction, or reflection. If it does none of these, remove it or move it behind a deliberate drill-down."),
        ("Actions alter the system", "Every meaningful action must change at least one state variable, constraint, relationship, or future option. Cosmetic actions weaken trust."),
        ("Consequences remain inspectable", "The experience should preserve before, decision, after, and explanation so outcomes can be reconstructed rather than merely announced."),
        ("Information is staged, not dumped", "The first view shows the situation; domain views expose causes; action views support intervention; history shows accumulation and interaction over time."),
        ("Prediction and discovery stay in tension", "Provide enough forecast to support intentional decisions, but keep some system relationships available for discovery through play."),
        ("The interface does not teach its own theory", "Apply sound representational and load-management principles without filling the participant experience with design commentary."),
    ]
    for title, text in commitments:
        add_definition(doc, title, text)
    add_callout(doc, "Removal rule", "When two areas present the same information, decide which one is for orientation and which one is for action. Merge or remove the competing surface.", fill=PALE_RED, accent="9B1C1C")


def add_model_skeleton(doc):
    add_heading(doc, "2. System model skeleton", 1)
    add_para(doc, "Build the model before polishing the interface. A thin but coherent model is more useful than a rich-looking interface whose parts do not affect one another.")
    add_heading(doc, "State specification", 2)
    for lead, text in (
        ("Outcome state", "What counts as progress, balance, failure, recovery, or completion?"),
        ("Operational state", "What is available, occupied, trained, assigned, waiting, blocked, or underused?"),
        ("Economic state", "What comes in, what goes out, which costs change with activity, and what creates a viability threshold?"),
        ("Human state", "What affects motivation, fatigue, confidence, trust, cooperation, retention, and usable effort?"),
        ("External state", "How do demand, willingness, reputation, referrals, policy, access, or environmental pressure change?"),
    ):
        add_bullet(doc, f"{lead}: {text}", bold_lead=f"{lead}:")
    add_heading(doc, "Actor and resource specification", 2)
    add_bullet(doc, "Separate roles when they contribute different kinds of capacity, authority, knowledge, or coordination.")
    add_bullet(doc, "Represent individual capability, available time, preferences, relationships, and load tolerance when these change outcomes.")
    add_bullet(doc, "Allow independent work, coordinated work, and pooled support when the real system contains all three patterns.")
    add_bullet(doc, "Translate percentages into concrete units: hours, cases, tasks, reach, cost, or another domain-relevant quantity.")
    add_heading(doc, "Constraint specification", 2)
    add_para(doc, "For every constraint, define the evidence that reveals it, the action that can address it, the delay before relief, and the new constraint likely to emerge afterward.")
    table = doc.add_table(rows=1, cols=4)
    headers = ["Constraint", "Visible evidence", "Available response", "Likely trade-off"]
    for i, label in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        set_run(p.add_run(label), size=9.5, color=NAVY, bold=True)
    for row in (
        ("Capacity", "Backlog, unmet demand, overload", "Reassign, add, train, reduce scope", "Cost, fatigue, opportunity loss"),
        ("Prerequisite", "Blocked option and missing requirement", "Acquire, prepare, partner, sequence", "Delay and fixed commitment"),
        ("Affordability", "Negative margin, weak runway, break-even gap", "Price, mix, cost, reach", "Demand or quality pressure"),
        ("Relationship", "Low trust, weak cooperation, reduced referrals", "Communicate, coordinate, invest", "Time and indirect payoff"),
    ):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_run(cells[i].paragraphs[0].add_run(value), size=9.3, color=INK)
        set_cant_split(table.rows[-1])
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [1500, 2460, 2460, 2940])


def add_choice_feedback(doc):
    add_heading(doc, "3. Choice, action, and consequence", 1)
    add_heading(doc, "Meaningful choice test", 2)
    add_bullet(doc, "At least two different strategies can plausibly reach the objective.")
    add_bullet(doc, "Options differ in timing, cost, risk, reversibility, or who benefits.")
    add_bullet(doc, "Participants can act, wait, reverse, or deliberately pass without being forced into a move.")
    add_bullet(doc, "Action limits are explicit and configurable when they are part of the challenge.")
    add_bullet(doc, "Blocked actions explain the missing condition and provide a direct route to address it.")
    add_heading(doc, "Action contract", 2)
    add_para(doc, "Specify every action using the same contract:")
    for text in (
        "Name: a plain-language command.",
        "Target: the state, actor, resource, or relationship being changed.",
        "Cost: money, time, attention, capacity, risk, or an action slot.",
        "Preconditions: what must already be true.",
        "Immediate effect: what changes now.",
        "Delayed effect: what changes after one or more cycles.",
        "Interaction effect: which other variables strengthen or weaken the result.",
        "Evidence: where the participant will see that the action worked.",
    ):
        add_bullet(doc, text)
    add_callout(doc, "Integrity rule", "Do not count only dramatic purchases or policy changes as actions. If reallocating people, changing schedules, adjusting price, or changing focus alters the model, it must be represented consistently in the action system.", fill=PALE_GOLD, accent="8A6500")
    add_heading(doc, "Consequence contract", 2)
    add_para(doc, "After each cycle, preserve a factual chain: decision taken -> state changed -> output changed -> trade-off created -> new constraint. Provide exact values where they support analysis and a short sentence that can be discussed or written down.")


def add_information_architecture(doc):
    add_heading(doc, "4. Information architecture skeleton", 1)
    add_para(doc, "Each screen should answer one question. Do not organize solely by the structure of the underlying data model.")
    add_heading(doc, "Recommended view hierarchy", 2)
    views = [
        ("Dashboard", "What needs attention now?", "Five to seven high-value signals, visual state, main constraint, goal progress."),
        ("Domain overview", "What is causing this result?", "Sorted rows, comparisons, capacity/demand bars, relationships, direct drill-down."),
        ("Action view", "What can be changed here?", "Editable controls, prerequisites, cost, scope, and queued action state."),
        ("Forecast", "What may this plan change?", "Baseline, plan, delta, constraint shift, and uncertainties; collapsible when not needed."),
        ("History/results", "What happened and why?", "Trends, action markers, grouped consequences, before/after values, reflection notes."),
    ]
    table = doc.add_table(rows=1, cols=3)
    for i, label in enumerate(("View", "Question", "Content job")):
        set_cell_shading(table.rows[0].cells[i], PALE_BLUE)
        set_run(table.rows[0].cells[i].paragraphs[0].add_run(label), size=9.5, color=NAVY, bold=True)
    for view, question, job in views:
        cells = table.add_row().cells
        set_run(cells[0].paragraphs[0].add_run(view), size=9.5, color=NAVY, bold=True)
        set_run(cells[1].paragraphs[0].add_run(question), size=9.5, color=INK)
        set_run(cells[2].paragraphs[0].add_run(job), size=9.5, color=INK)
        set_cant_split(table.rows[-1])
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [1500, 2880, 4980])
    add_heading(doc, "Split-screen pattern", 2)
    add_para(doc, "When comparison while acting is central, keep a compact situation view in sight and place interventions in an adjacent workspace. The persistent side should contain only information that is repeatedly consulted during decisions. Everything else belongs in tabs, drill-downs, or a collapsible drawer.")
    add_heading(doc, "Progressive nesting", 2)
    add_bullet(doc, "Overview: signal and ranking.")
    add_bullet(doc, "Selection: cause, requirement, and local relationship.")
    add_bullet(doc, "Action: editable control and immediate forecast.")
    add_bullet(doc, "Why: definition, unit, calculation basis, and deeper evidence.")
    add_bullet(doc, "History: accumulation and interaction over time.")


def add_representation(doc):
    add_heading(doc, "5. Representation and visual analysis", 1)
    add_para(doc, "Use multiple representations only when each performs a different job and participants can translate between them.")
    add_heading(doc, "Representation jobs", 2)
    add_definition(doc, "Orientation", "A small set of visual signals establishes the current situation and directs attention.")
    add_definition(doc, "Comparison", "Aligned rows, bars, or tables help participants compare alternatives using shared units.")
    add_definition(doc, "Explanation", "Requirements, causal links, and before/after values show why an outcome occurred.")
    add_definition(doc, "Action", "Controls sit beside the variable or blocker they change; the participant does not hunt for the fix.")
    add_definition(doc, "Accumulation", "Trend lines, stacked bars, event markers, and histories show change over time.")
    add_heading(doc, "Combined trend analysis", 2)
    add_bullet(doc, "Allow several measures to be shown together when their interaction is the analytic task.")
    add_bullet(doc, "Use a common time axis and label decision points directly on the chart.")
    add_bullet(doc, "Support baseline, planned, and actual series rather than a single unexplained line.")
    add_bullet(doc, "Keep exact values available through labels, tooltips, or a table adjacent to the visual.")
    add_bullet(doc, "Avoid combining measures whose scales make the comparison misleading; normalize or separate them clearly.")
    add_callout(doc, "Visual rule", "A graphic must help the participant notice, compare, predict, or explain. Decoration is not a substitute for analysis.", fill=PALE_TEAL, accent=TEAL)
    add_heading(doc, "Metric clarity card", 2)
    add_para(doc, "Every important term should have a compact definition containing: plain-language meaning, unit, time period, controllability, relevant threshold, source variables, and the actions that can change it.")


def add_load_clarity(doc):
    add_heading(doc, "6. Clarity and load management", 1)
    add_para(doc, "Complexity should come from the system and its trade-offs, not from decoding the interface.")
    for title, text in (
        ("Signal before detail", "Show the state and priority first; reveal formulas, assumptions, and full records on demand."),
        ("One question per view", "A screen that tries to orient, explain, edit, forecast, and document at once creates avoidable split attention."),
        ("Keep related items together", "Place blockers beside fixes, assignments beside their capacity meaning, and prices beside demand/revenue consequences."),
        ("Use consistent units", "Always state whether a value is per cycle, per person, per case, total, average, percentage, or remaining capacity."),
        ("Replace vague labels", "Terms such as readiness, sensitivity, impact, or coverage require an operational meaning that supports calculation or prediction."),
        ("Preserve position during edits", "Do not jump the user to the top, lose focus, swap row identity, or reorder controls unexpectedly after a change."),
        ("Control visual density", "Prefer sorted rows, compact bars, and progressive disclosure over repeated metric-card grids."),
        ("Reduce memory demands", "Use labeled states and progress chips; do not make participants remember what colors, dots, or abbreviations mean."),
    ):
        add_definition(doc, title, text)
    add_heading(doc, "Clarity test for any term", 2)
    for q in (
        "Can a participant define it in one sentence?",
        "Can they identify its unit and time basis?",
        "Can they name at least one cause and one action that changes it?",
        "Can they use it to make a prediction?",
        "Can they verify after a cycle whether their prediction was right?",
    ):
        add_bullet(doc, q)


def add_human_social(doc):
    add_heading(doc, "7. Human and social dynamics", 1)
    add_para(doc, "Human factors should modify the same operational model as money, capacity, demand, and facilities. They should not sit in a separate story layer with no consequence.")
    add_heading(doc, "Reusable human variables", 2)
    add_bullet(doc, "Capability and role boundaries: who can perform, authorize, support, or coordinate work.")
    add_bullet(doc, "Preferences and fit: which assignments improve or reduce usable effort.")
    add_bullet(doc, "Load and switching: how volume, fragmentation, and fatigue change effectiveness.")
    add_bullet(doc, "Pairing and networks: how relationships alter coordination, referrals, trust, or throughput.")
    add_bullet(doc, "Climate and confidence: how past outcomes affect future participation and performance.")
    add_bullet(doc, "Community or stakeholder pressure: how unmet need, access, fairness, or legitimacy changes demand and constraints.")
    add_heading(doc, "Transparency rule", 2)
    add_para(doc, "If motivation or social conditions change capacity, demand, or quality, show the adjustment, the direction, and a short reason. Avoid unexplained random penalties. Use randomness only when uncertainty is part of the intended challenge and the distribution is fair and reviewable.")


def add_trust_testing(doc):
    add_heading(doc, "8. Interaction trust and functional integrity", 1)
    add_para(doc, "Participants cannot learn from a model they do not trust. Small assignment or identity errors can invalidate later reasoning even when the underlying formula is sound.")
    add_heading(doc, "Trust requirements", 2)
    for item in (
        "Stable identity: an edited row, person, service, or option remains the same after re-rendering.",
        "Visible compatibility: incompatible choices may be allowed, but they must show zero useful contribution and the reason.",
        "No silent mutation: every model-changing edit appears in the plan and action count.",
        "Direct reachability: every visible blocker links to the exact action or prerequisite that can address it.",
        "Predictable persistence: focus, selection, open panel, and scroll position remain stable during fluid editing.",
        "Reversible planning: queued changes can be inspected and removed before the cycle is committed.",
        "Consistent accounting: the same action is not counted differently depending on the screen from which it was made.",
        "Graceful inaction: passing a cycle without a new action is allowed and still produces consequences from the current state.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "Full-path test", 2)
    add_step(doc, "Observe", "Identify a problem from the overview without prior explanation.")
    add_step(doc, "Diagnose", "Open the relevant domain and find the cause or missing requirement.")
    add_step(doc, "Act", "Make a change from the same context or a clearly linked action view.")
    add_step(doc, "Predict", "Read baseline, plan, delta, and possible constraint shift.")
    add_step(doc, "Commit", "Advance the cycle, including the option to commit no new action.")
    add_step(doc, "Explain", "Use the outcome record and trends to reconstruct what changed and why.")


def add_facilitation(doc):
    add_heading(doc, "9. Facilitation and debrief", 1)
    add_para(doc, "The system should produce material for facilitation rather than attempt to replace facilitation. The facilitator introduces the rules, observes reasoning, and supports reflection; the interactive model makes accumulation and constraint visible during play.")
    add_heading(doc, "Minimum debrief record", 2)
    add_bullet(doc, "The objective and starting condition.")
    add_bullet(doc, "Actions taken, including deliberate inaction.")
    add_bullet(doc, "Expected effects stated before commitment.")
    add_bullet(doc, "Actual changes in outputs, resources, relationships, and constraints.")
    add_bullet(doc, "A concise causal sentence generated from facts, not interpretation alone.")
    add_bullet(doc, "Space for the participant's explanation, uncertainty, and next hypothesis.")
    add_heading(doc, "Debrief prompts", 2)
    add_bullet(doc, "What did you think the main constraint was, and what evidence supported that view?")
    add_bullet(doc, "Which action produced an unexpected secondary effect?")
    add_bullet(doc, "What changed immediately, and what changed only after time passed?")
    add_bullet(doc, "Which measure improved while another deteriorated?")
    add_bullet(doc, "What would you keep, reverse, or test next?")


def add_workflow_checklist(doc):
    add_heading(doc, "10. Delivery workflow and acceptance gates", 1)
    add_heading(doc, "Build sequence", 2)
    for title, detail in (
        ("Model on paper", "Define state, actions, dependencies, and transition rules before interface polish."),
        ("Prove a viable path", "Demonstrate at least one winning or successful strategy and several plausible alternatives."),
        ("Create the smallest playable loop", "Observe -> decide -> forecast -> commit -> inspect consequence."),
        ("Add representations by job", "Orientation first, then diagnosis, action, comparison, and history."),
        ("Add human and social modifiers", "Only when they alter the core model and remain explainable."),
        ("Play through repeatedly", "Test inaction, edge cases, incompatible assignments, action limits, and recovery from poor decisions."),
        ("Test with target participants", "Observe where terms, relationships, or action routes require facilitator rescue."),
        ("Remove before adding", "Delete duplicate views, idle panels, long explanations, and metrics that do not guide a decision."),
    ):
        add_step(doc, title, detail)
    add_heading(doc, "Acceptance checklist", 2)
    checks = [
        "The objective is visible, measurable, and achievable with the starting system.",
        "More than one strategy can plausibly succeed.",
        "Every important scarcity produces visible evidence.",
        "Every visible blocker has a reachable response.",
        "Every action changes the model and appears in the plan.",
        "Participants may advance without taking a new action.",
        "The first view contains only high-priority signals.",
        "Each domain view answers one clear question.",
        "Terms include units, time basis, and practical meaning.",
        "Assignments and selections preserve identity after editing.",
        "Forecasts compare baseline, plan, and delta without revealing every answer.",
        "Results preserve concrete before/after consequences.",
        "Trends can combine relevant measures and mark decisions over time.",
        "Human and social variables affect the same operational model.",
        "The interface has no idle elements competing for attention.",
        "A full playthrough can be explained from the recorded evidence.",
    ]
    for item in checks:
        add_bullet(doc, item)


def add_canvas(doc):
    doc.add_page_break()
    add_heading(doc, "Reusable design canvas", 1)
    add_para(doc, "Complete this page before detailed interface design. Keep each answer short enough to expose uncertainty.")
    sections = [
        ("1. Objective", "What must participants achieve? What is the time horizon? What counts as success, failure, and recovery?"),
        ("2. Starting situation", "Which resources, capabilities, relationships, and constraints exist at the beginning?"),
        ("3. Discoverable scarcity", "What is missing or imbalanced? Which evidence allows participants to infer it?"),
        ("4. Choice space", "Which actions are available? What do they cost? Which paths and trade-offs differ meaningfully?"),
        ("5. Transition rules", "How does each action alter state now, later, and in combination with other variables?"),
        ("6. Representations", "Which view or visual supports orientation, diagnosis, action, prediction, comparison, and history?"),
        ("7. Human and social model", "Which motivations, relationships, trust conditions, or stakeholder pressures change outcomes?"),
        ("8. Consequence record", "What exact before/after values and factual sentence will be preserved after each cycle?"),
        ("9. Facilitation", "What must be introduced, observed, discussed, or written down outside the interface?"),
        ("10. Proof tests", "What will demonstrate playability, clarity, model trust, accessibility, and a viable success path?"),
    ]
    table = doc.add_table(rows=1, cols=2)
    for i, label in enumerate(("Design area", "Prompt")):
        set_cell_shading(table.rows[0].cells[i], PALE_BLUE)
        set_run(table.rows[0].cells[i].paragraphs[0].add_run(label), size=9.7, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    for label, prompt in sections:
        cells = table.add_row().cells
        set_cell_shading(cells[0], PALE_BLUE)
        p0 = cells[0].paragraphs[0]
        set_run(p0.add_run(label), size=9.7, color=NAVY, bold=True)
        p1 = cells[1].paragraphs[0]
        set_run(p1.add_run(prompt), size=9.7, color=INK)
        set_cant_split(table.rows[-1])
    set_table_geometry(table, [2000, 7360])
    add_callout(doc, "Final question", "Can a participant use the experience to form a hypothesis, make an intentional choice, observe a concrete consequence, and explain the relationship between the two?", fill=PALE_TEAL, accent=TEAL)


def build():
    doc = Document()
    style_document(doc)
    add_cover(doc)
    add_intro(doc)
    add_commitments(doc)
    add_model_skeleton(doc)
    add_choice_feedback(doc)
    add_information_architecture(doc)
    add_representation(doc)
    add_load_clarity(doc)
    add_human_social(doc)
    add_trust_testing(doc)
    add_facilitation(doc)
    add_workflow_checklist(doc)
    add_canvas(doc)
    doc.core_properties.title = "Interactive System Design Skeleton"
    doc.core_properties.subject = "Reusable guide for decision-centred simulations, sandboxes, and interactive learning tools"
    doc.core_properties.keywords = "interactive design, simulation design, decision sandbox, learning design, systems design"
    doc.core_properties.author = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
